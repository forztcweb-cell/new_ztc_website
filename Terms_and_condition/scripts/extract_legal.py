import json
import re
from pathlib import Path
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

SOURCE = Path('/home/ubuntu/upload/Zatics_Intelligence_Terms_and_Conditions_V2_Completed.docx')
OUTPUT = Path('/home/ubuntu/zaticsterms/public/legal-content.json')


def text_of_cell(cell):
    return '\n'.join(p.text.strip() for p in cell.paragraphs if p.text.strip()).strip()


def slug(value):
    value = re.sub(r'[^a-z0-9]+', '-', value.lower()).strip('-')
    return value[:90]


def split_numbered(text):
    text = text.replace('\u00a0', ' ').strip()
    match = re.match(r'^(\d+(?:\.\d+){0,2})[\.\t ]+(.*)$', text, re.S)
    if match:
        return match.group(1), match.group(2).strip()
    return None, text


def paragraph_block(p, index):
    raw = p.text.replace('\u00a0', ' ').strip()
    if not raw:
        return None
    num, rest = split_numbered(raw)
    style = p.style.name if p.style else ''
    if num and num.count('.') == 0 and rest and rest.upper() == rest:
        kind = 'h1'
        anchor = f'section-{num}-{slug(rest)}'
    elif num and num.count('.') == 1 and rest:
        kind = 'h2'
        anchor = f'clause-{num}-{slug(rest)}'
    elif num and num.count('.') == 2:
        kind = 'clause'
        anchor = f'clause-{num}'
    elif style.startswith('Heading 1'):
        kind = 'h1'; anchor = f'section-{slug(raw)}'
    elif style.startswith('Heading 2'):
        kind = 'h2'; anchor = f'clause-{slug(raw)}'
    else:
        kind = 'p'; anchor = None
    return {'type': kind, 'id': anchor, 'number': num, 'text': rest if num else raw, 'style': style, 'sourceIndex': index}


def table_block(table, index):
    rows = []
    for row in table.rows:
        rows.append([text_of_cell(cell) for cell in row.cells])
    return {'type': 'table', 'rows': rows, 'sourceIndex': index}


def main():
    doc = Document(str(SOURCE))
    # The first table is the document metadata table; keep it separate for the page header.
    metadata = {}
    if doc.tables:
        for row in doc.tables[0].rows:
            cells = [text_of_cell(c) for c in row.cells]
            if len(cells) >= 2 and cells[0]:
                metadata[cells[0]] = cells[1]

    blocks = []
    body_index = 0
    for child in doc.element.body.iterchildren():
        if child.tag == qn('w:p'):
            p = Paragraph(child, doc)
            block = paragraph_block(p, body_index)
            if block:
                blocks.append(block)
            body_index += 1
        elif child.tag == qn('w:tbl'):
            table = Table(child, doc)
            # Skip the cover metadata table and keep every legal-content table.
            if len(blocks) > 0:
                blocks.append(table_block(table, body_index))
            body_index += 1

    # Remove cover/title/TOC material. Legal content begins at the first numbered section.
    first_section = next((i for i, b in enumerate(blocks) if b.get('type') == 'h1' and b.get('number') == '1' and b.get('style', '').startswith('Heading')), 0)
    legal_blocks = blocks[first_section:]
    # De-duplicate generated anchors while preserving every block.
    used = set()
    for block in legal_blocks:
        if block.get('id'):
            base = block['id']; candidate = base; suffix = 2
            while candidate in used:
                candidate = f'{base}-{suffix}'; suffix += 1
            block['id'] = candidate; used.add(candidate)

    payload = {
        'metadata': metadata,
        'blocks': legal_blocks,
        'source': {
            'title': 'Zatics Intelligence — Terms and Conditions',
            'version': metadata.get('Version', '2.0'),
            'effectiveDate': metadata.get('Terms Effective Date', ''),
            'reference': metadata.get('Document reference', ''),
            'officialWebsite': metadata.get('Official website', 'https://zatics.tech')
        }
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f'wrote {OUTPUT} with {len(legal_blocks)} blocks')
    print('metadata:', metadata)
    print('h1:', sum(b['type'] == 'h1' for b in legal_blocks), 'h2:', sum(b['type'] == 'h2' for b in legal_blocks), 'tables:', sum(b['type'] == 'table' for b in legal_blocks))


if __name__ == '__main__':
    main()
